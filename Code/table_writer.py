"""Write portrait Word tables with their numbered captions and source notes."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from table_definitions import TABLES


def manuscript_text(value):
    """Use keyboard hyphens for ranges and compound terms; retain minus signs."""
    return str(value).replace('\u2013', '-')


def write_table(key, panels, folder):
    definition = TABLES[key]
    document = Document()
    section = document.sections[0]
    for name, value in definition['page'].items():
        setattr(section, name, Inches(value))
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1
    properties = document.core_properties
    properties.author = ''
    properties.last_modified_by = ''
    properties.comments = ''
    properties.title = manuscript_text(next(v for kind,v in definition['blocks'] if kind=='paragraph'))[:255]
    panel_index = 0
    for kind, block in definition['blocks']:
        if kind == 'paragraph':
            paragraph = document.add_paragraph(manuscript_text(block))
            paragraph.paragraph_format.widow_control = True
            if block.startswith(('Table ', 'A. ', 'B. ', '(a)', '(b)')):
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.bold = True
            continue
        rows = panels[panel_index]
        panel_index += 1
        headings = block['headers']
        if any(len(row) != len(headings) for row in rows):
            raise ValueError(f'{key}: table column count differs')
        table = document.add_table(rows=1, cols=len(headings))
        table.autofit = False
        table.style = 'Table Grid'
        for column, width in zip(table.columns, block['widths']):
            column.width = Inches(width)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, values in enumerate([headings]+rows):
            row = table.rows[0] if index==0 else table.add_row()
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            for cell, text, width in zip(row.cells, values, block['widths']):
                cell.width = Inches(width)
                cell.text = manuscript_text(text)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(block['font_size'])
                        run.bold = index == 0
                if index==0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'F2F2F2')
                    cell._tc.get_or_add_tcPr().append(shading)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
    if panel_index != len(panels):
        raise ValueError(f'{key}: unused table panel')
    destination = folder/(definition['stem']+'.docx')
    document.save(destination)
    return destination
